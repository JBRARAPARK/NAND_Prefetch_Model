"""Build a source-backed 20-page architecture PDF and editable Word file."""
import csv
from dataclasses import asdict
import hashlib
from html import escape
import json
import math
from pathlib import Path
import sys
import tempfile
import urllib.request

import fitz
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
ASSETS=OUT/'assets'
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(OUT))
from multiport_sim import MultiportConfig,MultiportSimulator
from content import build_sections

STEM='NAND_Prefetch_Architecture_Detailed_KO'
TITLE='NAND 프리페치 모델 전체 구조 설명'
MODEL_COMMIT='ee8636d53db93bcbe6a26025ef9cc0e939a55ac6'
NAVY='#18374D';TEAL='#147D83';GRAY='#586C7A';PALE='#EDF5F6';ORANGE='#D3893D'
WIDTH=507.3
FIG_HEIGHTS={'architecture':300,'demand_flow':212,'nand_states':185,'timing':160,
             'bandwidth':174,'decision_outstanding':174,'energy':174}


class EvidenceData:
    def __init__(self):
        self.commit=MODEL_COMMIT
        self.defaults=asdict(MultiportConfig())
        self.sources={'multiport':'results/multiport_256/sweep.csv',
                      'decision':'results/prefetch_decision_64/sweep.csv'}
        self.rows={k:list(csv.DictReader((ROOT/v).open())) for k,v in self.sources.items()}
        self.selected={}
        self.timing_cases=[]

    def select(self,group,**criteria):
        matches=[r for r in self.rows[group] if all(r[k]==str(v) for k,v in criteria.items())]
        if len(matches)!=1:raise ValueError((group,criteria,len(matches)))
        key=group+json.dumps(criteria,sort_keys=True)
        self.selected[key]={'source':self.sources[group],'criteria':criteria,'row':matches[0]}
        return matches[0]

    def multi(self,case,design='slow',variant='same_voltage_gated'):
        return self.select('multiport',case=case,design=design,power_variant=variant)

    def decision(self,case):return self.select('decision',case=case)


def fonts():
    folder=Path(tempfile.gettempdir())/'nand-executive-fonts'
    folder.mkdir(exist_ok=True)
    for weight,name in [('Regular','Korean'),('Bold','KoreanBold')]:
        path=folder/f'NanumGothic-{weight}.ttf'
        if not path.exists():
            urllib.request.urlretrieve(f'https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/NanumGothic-{weight}.ttf',path)
        pdfmetrics.registerFont(TTFont(name,str(path)))
        font_manager.fontManager.addfont(str(path))
    pdfmetrics.registerFontFamily('Korean',normal='Korean',bold='KoreanBold')
    plt.rcParams.update({'font.family':'NanumGothic','font.size':10,'axes.unicode_minus':False,
                         'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':GRAY,
                         'text.color':NAVY,'xtick.color':GRAY,'ytick.color':GRAY,
                         'svg.fonttype':'path','pdf.fonttype':42})


def box(c,x,y,w,h,title,lines=(),fill=PALE):
    c.setFillColor(colors.HexColor(fill));c.setStrokeColor(colors.HexColor('#C6D9DF'))
    c.roundRect(x,y,w,h,6,fill=1,stroke=1)
    c.setFillColor(colors.HexColor(NAVY));c.setFont('KoreanBold',9.4)
    c.drawCentredString(x+w/2,y+h-16,title)
    c.setFont('Korean',8.2)
    for i,line in enumerate(lines):c.drawCentredString(x+w/2,y+h-31-12*i,line)


def arrow(c,points,color=TEAL,label=None,label_at=None):
    c.setStrokeColor(colors.HexColor(color));c.setLineWidth(1.35)
    path=c.beginPath();path.moveTo(*points[0])
    for point in points[1:]:path.lineTo(*point)
    c.drawPath(path)
    a,b=points[-2:];angle=math.atan2(b[1]-a[1],b[0]-a[0])
    for side in [-1,1]:
        c.line(b[0],b[1],b[0]-5*math.cos(angle+side*.5),b[1]-5*math.sin(angle+side*.5))
    if label:
        c.setFillColor(colors.HexColor(color));c.setFont('Korean',7.6)
        c.drawString(*label_at,label)


def architecture(c):
    for port in range(4):
        x=port*129.1
        box(c,x,205,120,98,f'Host Port {port}',[
            '512bit · 1GHz · 64GB/s','AR + outstanding','Host 판단 → 조회 FIFO','2 lanes / 500 slots','포트 FIFO → R 반환'])
        arrow(c,[(x+51,205),(x+51,183)])
        arrow(c,[(x+68,183),(x+68,205)],color=ORANGE)
    box(c,0,147,507,36,'공유 page 상태 · 동일 page 읽기 병합',[])
    box(c,0,80,208,36,'예측 → 선행 판단 → page 요청',[])
    box(c,0,0,242,57,'공유 Controller Buffer',[
        '8MiB · page credit 선점','page 전체 burst 반환 후 공간 해제'])
    box(c,261,0,246,57,'공유 NAND Scheduler',[
        '128 chips × 6 planes / 128 channels','command → sense → transfer → ECC'])
    arrow(c,[(208,98),(383,98),(383,57)],label='speculative issue',label_at=(240,102))
    arrow(c,[(475,147),(475,57)],label='demand issue',label_at=(391,126))
    arrow(c,[(261,24),(242,24)],color=ORANGE)
    arrow(c,[(229,57),(229,147)],color=ORANGE,label='ready 데이터',label_at=(235,70))
    c.setFont('Korean',7.2);c.setFillColor(colors.HexColor(GRAY))
    c.drawString(0,128,'청록: 요청·제어     주황: 데이터·ready 알림     실제 fabric 중재는 별도 가정')


def demand_flow(c):
    for x,title,lines in [(0,'AR 수락',['Outstanding +1']),
                          (178,'Host 판단 (선택)',['100 / 200 / 300ns']),
                          (356,'Lookup FIFO / engine',['슬롯·lane·clock edge'])]:
        box(c,x,165,150,45,title,lines)
    arrow(c,[(150,187),(178,187)]);arrow(c,[(328,187),(356,187)])
    box(c,356,95,150,45,'Page 상태 확인',['ready / late / miss'])
    box(c,178,95,150,45,'NAND 읽기 / 병합',['miss 생성 · late 대기'])
    box(c,0,95,150,45,'Ready data',['lookup + page 준비'])
    arrow(c,[(431,165),(431,140)])
    arrow(c,[(356,116),(328,116)],label='miss/late',label_at=(321,143))
    arrow(c,[(178,116),(150,116)],color=ORANGE)
    arrow(c,[(431,95),(431,75),(75,75),(75,95)],color=ORANGE,label='ready hit',label_at=(230,80))
    for x,title,lines in [(0,'포트 내부 FIFO',['선두 burst만 반환']),
                          (178,'R 데이터 전송',['64B · 1ns']),
                          (356,'RLAST / 공간 해제',['O -1; 전체 page면 credit 해제'])]:
        box(c,x,0,150,45,title,lines)
    arrow(c,[(75,95),(75,45)],color=ORANGE)
    arrow(c,[(150,22),(178,22)],color=ORANGE);arrow(c,[(328,22),(356,22)],color=ORANGE)


def nand_states(c):
    upper=[('Queued','자원·credit 대기'),('Command','chip 점유 0.1µs'),('Sense','tR 2.08µs'),('Sensed','channel 대기')]
    lower=[('Returned','page credit 해제'),('Ready','burst 반환 가능'),('ECC','0.2µs'),('Transfer','6.827 + 0.05µs')]
    for row,y in [(upper,125),(lower,48)]:
        for i,(title,line) in enumerate(row):box(c,i*130, y,117,48,title,[line])
    for i in range(3):arrow(c,[(i*130+117,149),((i+1)*130,149)])
    arrow(c,[(448,125),(448,96)])
    for i in range(3,0,-1):arrow(c,[(i*130,72),((i-1)*130+117,72)],color=ORANGE)
    c.setFont('Korean',8);c.setFillColor(colors.HexColor(GRAY))
    c.drawString(0,105,'Sense 시작: command 해제      Transfer 완료: channel·plane 해제')
    c.drawString(0,26,'Buffer credit: Command 직전 선점 → page의 모든 burst가 반환된 뒤 해제')
    c.drawString(0,10,'같은 page를 여러 포트가 참조해도 NAND read·buffer credit은 한 번만 사용')


def timing(c,d):
    cases=[('Ready · 추가 판단 없음',{}),
           ('AR 이후 판단 300ns',{'host_decision_us':.3}),
           ('선행 판단 300ns · 짧은 lead',{'prefetch_decision_us':.3,'early_lead_us':8.756666666666666})]
    origin=172;scale=.30
    c.setFont('Korean',7.4);c.setFillColor(colors.HexColor(GRAY))
    for j,(label,color) in enumerate([('판단',ORANGE),('조회',TEAL),('NAND 잔여 대기','#A6B6C2')]):
        c.setFillColor(colors.HexColor(color));c.rect(j*123,170,10,8,fill=1,stroke=0)
        c.setFillColor(colors.HexColor(GRAY));c.drawString(j*123+15,171,label)
    for i,(label,override) in enumerate(cases):
        config=MultiportConfig(requests=1,request_bytes=16384,scenario='ready',outstanding=128,
                              record_trace=False,**override)
        sim=MultiportSimulator(config);sim.run();burst=sim.bursts[0];epoch=config.first_demand_us
        relative={k:(burst[k]-epoch)*1000 for k in ['ar','decision_complete_us','lookup_start','lookup_complete_us','r_start','rlast']}
        d.timing_cases.append({'label':label,'config':asdict(config),'first_burst_ns':relative,
                               'page_ready_ns':(sim.pages[0]['ready']-epoch)*1000})
        y=132-i*42
        c.setFillColor(colors.HexColor(NAVY));c.setFont('Korean',8)
        c.drawString(0,y+3,label)
        for start,end,color in [(relative['ar'],relative['decision_complete_us'],ORANGE),
                                (relative['lookup_start'],relative['lookup_complete_us'],TEAL),
                                (relative['lookup_complete_us'],relative['r_start'],'#A6B6C2')]:
            if end>start+.01:
                c.setFillColor(colors.HexColor(color));c.rect(origin+start*scale,y,(end-start)*scale,13,fill=1,stroke=0)
        x=origin+relative['rlast']*scale
        c.setFillColor(colors.HexColor(NAVY));c.circle(x,y+6.5,2.2,fill=1,stroke=0)
        c.setFont('KoreanBold',8);c.drawString(x+5,y+3,f'{relative["rlast"]:.0f}ns')
    c.setStrokeColor(colors.HexColor('#9AAEBB'));c.line(origin,27,origin+1000*scale,27)
    c.setFont('Korean',7.3);c.setFillColor(colors.HexColor(GRAY))
    for ns in [0,200,400,600,800,1000]:
        x=origin+ns*scale;c.line(x,27,x,23);c.drawCentredString(x,13,str(ns))
    c.drawString(0,13,'첫 AR 이후 상대 시간 / ns')


def diagrams(d):
    for name,height,draw in [('architecture',310,architecture),('demand_flow',214,demand_flow),('nand_states',185,nand_states)]:
        path=ASSETS/f'{name}.pdf'
        c=canvas.Canvas(str(path),pagesize=(527,height+20))
        c.translate(10,10);draw(c);c.save()
        with fitz.open(path) as pdf:pdf[0].get_pixmap(matrix=fitz.Matrix(2.2,2.2),alpha=False).save(ASSETS/f'{name}.png')
    path=ASSETS/'timing.pdf';c=canvas.Canvas(str(path),pagesize=(527,205))
    c.translate(10,10);timing(c,d);c.save()
    with fitz.open(path) as pdf:pdf[0].get_pixmap(matrix=fitz.Matrix(2.2,2.2),alpha=False).save(ASSETS/'timing.png')


def save_chart(fig,name):
    fig.tight_layout(pad=.8)
    fig.savefig(ASSETS/f'{name}.png',dpi=180,bbox_inches='tight',facecolor='white')
    fig.savefig(ASSETS/f'{name}.svg',bbox_inches='tight',facecolor='white')
    plt.close(fig)


def charts(d):
    fig,ax=plt.subplots(figsize=(7.2,2.6))
    channels=[4,32,64,107,128];positions=list(range(5))
    rows=[d.multi(f'nand_{n}ch') for n in channels]
    ax.bar([x-.18 for x in positions],[float(r['throughput_gbps']) for r in rows],width=.35,color=TEAL,label='전체 평균')
    ax.bar([x+.18 for x in positions],[float(r['tail_throughput_gbps']) for r in rows],width=.35,color=ORANGE,label='후반 평균')
    ax.axhline(256,color=NAVY,linewidth=1,linestyle='--');ax.text(-.45,263,'Host 상한 256GB/s',fontsize=9)
    ax.set_xticks(positions,[f'{n}채널' for n in channels]);ax.set_ylabel('4포트 합계 GB/s');ax.set_ylim(0,305)
    ax.legend(frameon=False,loc='upper left');ax.grid(axis='y',alpha=.12);ax.set_axisbelow(True)
    save_chart(fig,'bandwidth')

    fig,ax=plt.subplots(figsize=(7.2,2.6));ns=[0,100,200,300]
    ax.plot(ns,[128]*4,marker='o',color=TEAL,linewidth=2,label='선행 판단 · lead 10µs')
    ax.plot(ns,[128,160,192,224],marker='s',color=ORANGE,linewidth=2,label='AR 이후 판단')
    for x,y in zip(ns,[128,160,192,224]):ax.annotate(str(y),(x,y),xytext=(0,8),textcoords='offset points',ha='center',fontsize=9)
    ax.set_xticks(ns);ax.set_xlabel('추가 판단 지연 ns');ax.set_ylabel('권장 outstanding / 포트');ax.set_ylim(110,250)
    ax.legend(frameon=False,loc='upper left');ax.grid(alpha=.15);save_chart(fig,'decision_outstanding')

    fig,ax=plt.subplots(figsize=(7.2,2.6));same=[];low=[]
    for case in ['ready_O502','nand_long']:
        fast=d.multi(case,'fast')
        same.append(100*float(d.multi(case)['total_energy_eu'])/float(fast['total_energy_eu']))
        low.append(100*float(d.multi(case,'slow','lower_voltage_gated')['total_energy_eu'])/float(fast['total_energy_eu']))
    for offsets,values,col,label in [(-.18,same,TEAL,'500MHz · 1V'),(.18,low,ORANGE,'500MHz · 0.8V 가정')]:
        bars=ax.bar([x+offsets for x in range(2)],values,width=.35,color=col,label=label)
        ax.bar_label(bars,labels=[f'{v:.1f}%' for v in values],padding=3,fontsize=9)
    ax.axhline(100,color=NAVY,linestyle='--',linewidth=1);ax.set_ylim(0,128)
    ax.set_xticks([0,1],['Ready · 8MiB','NAND · 32MiB']);ax.set_ylabel('동일 작업 총 에너지비 %')
    ax.legend(frameon=False,loc='upper right');ax.grid(axis='y',alpha=.12);ax.set_axisbelow(True)
    save_chart(fig,'energy')


def safe(text):return escape(str(text)).replace('\n','<br/>')


def build_pdf(pages):
    styles={
        'title':ParagraphStyle('title',fontName='KoreanBold',fontSize=21,leading=28,textColor=colors.HexColor(NAVY),spaceAfter=10,wordWrap='CJK'),
        'lead':ParagraphStyle('lead',fontName='KoreanBold',fontSize=10.8,leading=17,textColor=colors.HexColor(TEAL),spaceAfter=12,wordWrap='CJK'),
        'body':ParagraphStyle('body',fontName='Korean',fontSize=9.8,leading=15.6,textColor=colors.HexColor(NAVY),spaceAfter=8,wordWrap='CJK'),
        'heading':ParagraphStyle('heading',fontName='KoreanBold',fontSize=10.3,leading=16,textColor=colors.HexColor(TEAL),spaceBefore=6,spaceAfter=6,wordWrap='CJK'),
        'formula':ParagraphStyle('formula',fontName='KoreanBold',fontSize=9.1,leading=15,textColor=colors.HexColor(NAVY),wordWrap='CJK'),
        'cell':ParagraphStyle('cell',fontName='Korean',fontSize=8.1,leading=12,textColor=colors.HexColor(NAVY),wordWrap='CJK'),
        'th':ParagraphStyle('th',fontName='KoreanBold',fontSize=8.2,leading=12,textColor=colors.white,wordWrap='CJK'),
        'caption':ParagraphStyle('caption',fontName='Korean',fontSize=7.8,leading=11.7,textColor=colors.HexColor(GRAY),spaceAfter=8,wordWrap='CJK')}
    def para(text,kind='body'):return Paragraph(safe(text),styles[kind])
    story=[]
    for index,page in enumerate(pages):
        if index:story.append(PageBreak())
        story += [para(f'{index+1:02d}  {page["title"]}','title'),para(page['lead'],'lead')]
        for block in page['blocks']:
            kind=block['type']
            if kind in ['paragraph','heading']:story.append(para(block['text'],'body' if kind=='paragraph' else 'heading'))
            elif kind=='bullets':
                for item in block['items']:story.append(para('• '+item))
            elif kind=='formula':
                table=Table([[para(block['text'],'formula')]],colWidths=[WIDTH])
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor(PALE)),
                                          ('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10),
                                          ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
                story.extend([table,Spacer(1,8)])
            elif kind=='table':
                widths=block['widths'] or [100/len(block['headers'])]*len(block['headers'])
                rows=[[para(x,'th') for x in block['headers']]]+[[para(x,'cell') for x in row] for row in block['rows']]
                table=Table(rows,colWidths=[WIDTH*x/100 for x in widths],repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(NAVY)),
                                          ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F4F7F8')]),
                                          ('LINEBELOW',(0,0),(-1,-1),.35,colors.HexColor('#D5E1E6')),
                                          ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),
                                          ('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),
                                          ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
                story.extend([table,Spacer(1,9)])
            elif kind=='figure':
                story.append(Image(str(ASSETS/f'{block["name"]}.png'),width=WIDTH,height=FIG_HEIGHTS[block['name']],kind='proportional'))
                story.append(para(block['caption'],'caption'))
    def footer(c,doc):
        c.saveState();c.setFillColor(colors.HexColor(GRAY));c.setFont('Korean',7.5)
        c.drawString(44,814,'NAND Prefetch Model  |  구조 · 성능 · Outstanding · 파워')
        c.setStrokeColor(colors.HexColor('#D5E1E6'));c.line(44,802,551,802)
        c.drawString(44,25,'가정 기반 모델 · 2026-10-04 · source '+MODEL_COMMIT[:12])
        c.linkURL('https://github.com/JBRARAPARK/NAND_Prefetch_Model/tree/'+MODEL_COMMIT,
                  (44,20,380,34),relative=0,thickness=0)
        c.drawRightString(551,25,f'{doc.page} / {len(pages)}');c.restoreState()
    path=OUT/f'{STEM}.pdf'
    class DetailedTemplate(SimpleDocTemplate):
        def afterFlowable(self,flowable):
            if isinstance(flowable,Paragraph) and flowable.style.name=='title':
                key=f'chapter_{self.page}'
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.getPlainText(),key,level=0)
    doc=DetailedTemplate(str(path),pagesize=(595.276,841.89),leftMargin=44,rightMargin=44,
                          topMargin=53,bottomMargin=43,title=TITLE,author='NAND Prefetch Model',
                          subject='4port 512bit 256GB/s, 64GB/s outstanding, decision latency and energy')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return path


def shade(cell,fill):
    element=OxmlElement('w:shd');element.set(qn('w:fill'),fill)
    cell._tc.get_or_add_tcPr().append(element)


def build_docx(pages):
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
    section.left_margin=section.right_margin=Inches(.61)
    section.top_margin=Inches(.62);section.bottom_margin=Inches(.60)
    for name,size in [('Normal',9.5),('Title',20),('Heading 1',10.5),('Caption',8)]:
        style=doc.styles[name];style.font.name='맑은 고딕';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string(NAVY[1:])
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.12
    doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string(TEAL[1:])
    header=section.header.paragraphs[0];header.text='NAND Prefetch Model | 전체 구조 상세 설명'
    for run in header.runs:run.font.size=Pt(8)
    foot=section.footer.paragraphs[0];foot.text='2026-10-04  ·  source '+MODEL_COMMIT[:12]+'  ·  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for index,page in enumerate(pages):
        if index:doc.add_page_break()
        doc.add_paragraph(f'{index+1:02d}  {page["title"]}','Title')
        lead=doc.add_paragraph(page['lead'])
        for run in lead.runs:run.bold=True;run.font.color.rgb=RGBColor.from_string(TEAL[1:])
        for block in page['blocks']:
            kind=block['type']
            if kind=='paragraph':doc.add_paragraph(block['text'])
            elif kind=='heading':doc.add_paragraph(block['text'],'Heading 1')
            elif kind=='bullets':
                for item in block['items']:doc.add_paragraph('• '+item)
            elif kind=='formula':
                p=doc.add_paragraph(block['text'])
                for run in p.runs:run.bold=True
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),PALE[1:]);p._p.get_or_add_pPr().append(shd)
            elif kind=='figure':
                doc.add_picture(str(ASSETS/f'{block["name"]}.png'),width=Inches(7.0))
                doc.add_paragraph(block['caption'],'Caption')
            elif kind=='table':
                table=doc.add_table(rows=1,cols=len(block['headers']));table.autofit=False
                widths=block['widths'] or [100/len(block['headers'])]*len(block['headers'])
                for i,label in enumerate(block['headers']):table.rows[0].cells[i].text=label
                for values in block['rows']:
                    cells=table.add_row().cells
                    for i,value in enumerate(values):cells[i].text=str(value)
                for ri,row in enumerate(table.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    for i,cell in enumerate(row.cells):
                        cell.width=Inches(7.0*widths[i]/100)
                        shade(cell,NAVY[1:] if ri==0 else ('F4F7F8' if ri%2==0 else 'FFFFFF'))
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
                            p.paragraph_format.line_spacing=1.08
                            for run in p.runs:
                                run.font.size=Pt(8)
                                if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                doc.add_paragraph('')
    doc.core_properties.title=TITLE
    doc.core_properties.subject='Host·프리페치·판단·조회·NAND·버퍼·반환·성능·파워 상세 설계 설명'
    path=OUT/f'{STEM}.docx';doc.save(path)
    return path


def validate(pages,pdfpath,docpath):
    with fitz.open(pdfpath) as pdf:
        assert len(pdf)==len(pages),(len(pdf),len(pages))
        for i,page in enumerate(pages):
            text=pdf[i].get_text()
            assert page['title'] in text,(i,page['title'])
            assert len(text)>300,(i,'Unexpectedly empty page')
            # PDF words must stay inside printable page bounds.
            for x0,y0,x1,y1,*rest in pdf[i].get_text('words'):
                assert x0>=20 and x1<=575 and y0>=15 and y1<=832,(i,rest)
        full='\n'.join(p.get_text() for p in pdf)
        for token in ['502','2,560','10,240','100ns','300ns','76.3%','55.8%','50개 테스트']:
            assert token in full,token
        # A contact sheet previews all 20 pages; generated separately from source.
        thumbs=[]
        for page in pdf:
            pix=page.get_pixmap(matrix=fitz.Matrix(.34,.34),alpha=False)
            thumbs.append(pix)
        from PIL import Image as PILImage,ImageDraw
        w,h=thumbs[0].width,thumbs[0].height
        sheet=PILImage.new('RGB',(w*4+50,h*5+120),'#DCE5EB')
        draw=ImageDraw.Draw(sheet)
        for i,pix in enumerate(thumbs):
            x=10+(i%4)*(w+10);y=10+(i//4)*(h+22)
            sheet.paste(PILImage.frombytes('RGB',[pix.width,pix.height],pix.samples),(x,y))
            draw.text((x,y+h+3),f'{i+1:02d}',fill='#18374D')
        sheet.save(OUT/'preview_all_pages.png')
        for i,name in [(1,'preview_structure.png'),(12,'preview_decision.png')]:
            pdf[i].get_pixmap(matrix=fitz.Matrix(1.3,1.3),alpha=False).save(OUT/name)
    doc=Document(docpath)
    alltext='\n'.join(p.text for p in doc.paragraphs)
    for page in pages:assert page['title'] in alltext,page['title']
    expected=[block for page in pages for block in page['blocks'] if block['type']=='table']
    assert len(doc.tables)==len(expected)
    for table,block in zip(doc.tables,expected):
        assert [[c.text for c in row.cells] for row in table.rows]==[block['headers']]+[[str(x) for x in row] for row in block['rows']]
    figures=sum(block['type']=='figure' for page in pages for block in page['blocks'])
    assert len(doc.inline_shapes)==figures
    return {'pdf_pages':len(pages),'pdf_chapter_titles_verified':len(pages),
            'pdf_text_inside_bounds':True,'word_tables_verified':len(expected),
            'word_figures_verified':figures,'word_note':'Explicit 20-section page breaks; pagination may vary by Word font and renderer.'}


def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    OUT.mkdir(exist_ok=True);ASSETS.mkdir(exist_ok=True)
    fonts();d=EvidenceData();pages=build_sections(d)
    diagrams(d);charts(d)
    content={'title':TITLE,'model_commit':MODEL_COMMIT,'pages':pages}
    contentpath=OUT/'content.json';contentpath.write_text(json.dumps(content,ensure_ascii=False,indent=2)+'\n')
    pdf=build_pdf(pages);word=build_docx(pages);validation=validate(pages,pdf,word)
    sources=['multiport_sim.py','axi_burst_sim.py','lookup_power.py','sim.py',
             'results/multiport_256/sweep.csv','results/multiport_256/manifest.json',
             'results/prefetch_decision_64/sweep.csv','results/prefetch_decision_64/configs.json',
             'results/prefetch_decision_64/manifest.json','test_multiport.py','test_prefetch_decision.py',
             'docs/architecture_details/build_report.py','docs/architecture_details/content.py']
    evidence={'model_commit':MODEL_COMMIT,'source_sha256':{name:digest(ROOT/name) for name in sources},
              'selected_source_rows':list(d.selected.values()),'validation':validation,
              'timing_cases':d.timing_cases,
              'output_sha256':{p.name:digest(p) for p in [pdf,word,contentpath]},
              'semantics':'20-page PDF. Word uses the same paragraphs/tables/figures. Historical runs remain source snapshots; no simulator or power behavior changed. AR-rate limiter, decision capacity and predictor accuracy are described as future work.'}
    (OUT/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(validation,ensure_ascii=False))
    print(pdf);print(word)


if __name__=='__main__':main()
