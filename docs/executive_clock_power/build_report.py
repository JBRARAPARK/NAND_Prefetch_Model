"""Create a one-page Korean executive PDF and an editable Word companion."""
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import urllib.request

import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
CSV=ROOT/'results/lookup_clock_power_64b/sweep.csv'
MANIFEST=ROOT/'results/lookup_clock_power_64b/manifest.json'
STEM='NAND_Lookup_Clock_Power_Executive_Brief_KO'
TITLE='NAND 프리페치 조회 클럭·에너지 검토'
SUBTITLE='임원 보고  |  2026년 10월 3일  |  모델 기반 설계 검토'
LEAD='조회 용량 확보가 성능을 지키고, 전압 저감이 에너지 절감 폭을 키웁니다.'
SUMMARY='1GHz·조회 0ns와 500MHz·조회 500ns를 동일 병렬성으로 비교했습니다. 저클럭의 이점은 조회 처리능력과 클럭 유휴 제어에 따라 달라집니다. 판단 기준은 평균 파워와 동일 작업량의 총 에너지를 함께 보는 것입니다.'
HEADERS=['500MHz 설계 조건','처리량\nGB/s','평균 파워비','총 에너지비']
TABLE_NOTE='비율 기준: 각 행과 같은 조회 슬롯 수의 1GHz·0ns·1V 설계 = 100%. 양쪽 모두 clock gating(유휴 시 클럭 차단) 적용. 1GHz 기준 처리량은 32.00GB/s입니다.'
TERMS='조회 슬롯은 동시에 판정 중인 요청의 수입니다. 파워는 단위시간당 소비, 에너지는 같은 데이터량의 처리에 든 총소비입니다. 에너지 = 평균 파워 × 실행시간.'
INSIGHT='슬롯 64개에서는 평균 파워가 64.9% 낮아도 실행시간이 3.91배로 늘어, 총 에너지는 37.3% 증가합니다.'
DECISIONS=[
    ('성능 확보', '32GB/s 목표라면 500ns 조회에서 슬롯 250개 이상, 미완료 요청 수(outstanding) 251개 이상을 함께 확보합니다. 파이프라인당 매 클럭 새 조회를 수락하는 조건입니다.'),
    ('절전 판단', '슬롯 250개·동일 전압·gating에서는 성능과 에너지가 거의 같습니다. 0.8V에서도 500MHz가 가능하면 모델 에너지는 26.3% 감소합니다. 전압·주파수 가능 범위와 실제 에너지 계수를 측정해 판단합니다.'),
    ('구현 우선순위', '조회 지연과 별도로 새 요청 수락 간격을 확인합니다. 단일 파이프라인이 2클럭마다 수락하면 500MHz 처리량 상한은 16GB/s로 낮아져, 슬롯 확대만으로는 부족합니다.')]
GATING='클럭 유휴 제어의 영향: 슬롯 250개·동일 전압에서 클럭을 항상 켜는 모델은 에너지가 9.6% 감소하지만, 양쪽에 gating을 적용하면 절감 이점이 거의 사라집니다.'
SCOPE='비교 범위: 데이터가 사전 준비된 ready hit, 64B burst, 32×64KiB 요청, AXI 256bit·1GHz 고정, outstanding 500, 독립 조회 파이프라인 1개·투입 간격 1클럭. 31.76GB/s는 초기 조회 지연을 포함한 평균입니다.'
LIMIT='파워·에너지는 예시 EU 계수의 상대 결과이며 제품 실측 W·J가 아닙니다. 0ns는 판정 지연을 생략한 이상적 기준입니다. 0.8V는 동작 가능성을 가정하며, 사전 prefetch 시간·에너지는 제외합니다.'
VALIDATION='검증: 테스트 35개 통과 · 성능 실행 216개 · 에너지 계산 648개 · 기존 32개 조건 회귀 일치.'
SOURCE='근거: results/lookup_clock_power_64b/sweep.csv 및 manifest.json'
PDF_CHECKS=['조회 클럭','핵심 비교 결과','137.3%','100.1%','73.7%','31.76','26.3%','16GB/s']
TABLE_HIGHLIGHT_ROW=3


def get_rows():
    rows=list(csv.DictReader(CSV.open()))
    def select(design,c,variant):
        return next(r for r in rows if r['group']=='main' and r['scenario']=='ready'
                    and r['outstanding']=='500' and r['lookup_slots']==str(c)
                    and r['design']==design and r['power_variant']==variant)
    table=[]
    for c,variant in [(64,'same_voltage_gated'),(250,'same_voltage_gated'),(250,'lower_voltage_gated')]:
        fast=select('fast',c,'same_voltage_gated');slow=select('slow',c,variant)
        table.append([f'슬롯 {c}개 · {float(slow["lookup_voltage_v"]):g}V',
                      f'{float(slow["throughput_gbps"]):.2f}',
                      f'{100*float(slow["mean_total_power_eu_per_us"])/float(fast["mean_total_power_eu_per_us"]):.1f}%',
                      f'{100*float(slow["total_energy_eu"])/float(fast["total_energy_eu"]):.1f}%'])
    return table


def register_fonts():
    fontdir=Path(tempfile.gettempdir())/'nand-executive-fonts';fontdir.mkdir(exist_ok=True)
    for weight,name in [('Regular','Korean'),('Bold','KoreanBold')]:
        path=fontdir/f'NanumGothic-{weight}.ttf'
        if not path.exists():
            urllib.request.urlretrieve(f'https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/NanumGothic-{weight}.ttf',path)
        pdfmetrics.registerFont(TTFont(name,str(path)))
    pdfmetrics.registerFontFamily('Korean',normal='Korean',bold='KoreanBold')


def build_pdf(rows):
    register_fonts()
    navy=colors.HexColor('#173750');teal=colors.HexColor('#147B80')
    styles={
        'body':ParagraphStyle('body',fontName='Korean',fontSize=9.4,leading=14.2,textColor=navy,spaceAfter=6,wordWrap='CJK'),
        'small':ParagraphStyle('small',fontName='Korean',fontSize=8,leading=11.6,textColor=colors.HexColor('#536779'),spaceAfter=5,wordWrap='CJK'),
        'title':ParagraphStyle('title',fontName='KoreanBold',fontSize=19,leading=25,textColor=navy,spaceAfter=12),
        'head':ParagraphStyle('head',fontName='KoreanBold',fontSize=10.5,leading=15,textColor=teal,spaceBefore=8,spaceAfter=5),
        'lead':ParagraphStyle('lead',fontName='KoreanBold',fontSize=11.2,leading=17,textColor=navy,wordWrap='CJK'),
        'cell':ParagraphStyle('cell',fontName='Korean',fontSize=10,leading=14.5,alignment=1,textColor=navy,wordWrap='CJK'),
        'th':ParagraphStyle('th',fontName='KoreanBold',fontSize=9,leading=13,alignment=1,textColor=colors.white,wordWrap='CJK')}
    def p(text,kind='body'):return Paragraph(text.replace('\n','<br/>'),styles[kind])
    story=[p(SUBTITLE,'small'),Spacer(1,4),p(TITLE,'title')]
    callout=Table([[p(LEAD,'lead')]],colWidths=[523])
    callout.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E8F3F4')),
                                ('LEFTPADDING',(0,0),(-1,-1),11),('RIGHTPADDING',(0,0),(-1,-1),11),
                                ('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]))
    story += [callout,Spacer(1,9),p(SUMMARY),p('핵심 비교 결과','head')]
    table=Table([[p(x,'th') for x in HEADERS]]+[[p(x,'cell') for x in row] for row in rows],
                colWidths=[181,104,119,119])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),
                              ('BACKGROUND',(0,TABLE_HIGHLIGHT_ROW),(-1,TABLE_HIGHLIGHT_ROW),colors.HexColor('#EDF6F4')),
                              ('LINEBELOW',(0,1),(-1,-1),.4,colors.HexColor('#D7E1E9')),
                              ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),
                              ('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    story += [table,Spacer(1,7),p(TABLE_NOTE,'small'),p(TERMS,'small'),p('<b>'+INSIGHT+'</b>'),
              p('설계 의사결정 제안','head')]
    for label,body in DECISIONS:story.append(p('<b>'+label+'</b>　'+body))
    story += [p(GATING,'small'),p('판단 범위와 근거','head'),p(SCOPE,'small'),p(LIMIT,'small'),
              p(VALIDATION,'small'),p(SOURCE,'small')]
    path=OUT/f'{STEM}.pdf'
    doc=SimpleDocTemplate(str(path),pagesize=(595.276,841.89),leftMargin=36,rightMargin=36,
                          topMargin=30,bottomMargin=30,title=TITLE,author='NAND Prefetch Model')
    doc.build(story)
    with fitz.open(path) as pdf:
        assert len(pdf)==1,f'Expected 1 page, got {len(pdf)}'
        text=pdf[0].get_text()
        for expected in PDF_CHECKS:
            assert expected in text,expected
        pdf[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(OUT/'preview.png')
    print('PDF: 1 page, Korean text and CSV-derived figures verified')


def shade(cell,fill):
    props=cell._tc.get_or_add_tcPr();x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);props.append(x)


def build_docx(rows):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.45)
    sec.left_margin=sec.right_margin=Inches(.5)
    normal=doc.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(9)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.05
    for name,size in [('Title',18),('Heading 1',10.5)]:
        st=doc.styles[name];st.font.name='맑은 고딕';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('173750')
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5)
    def paragraph(text,small=False,bold=False):
        p=doc.add_paragraph(text)
        for r in p.runs:
            r.bold=bold
            if small:r.font.size=Pt(8)
        return p
    paragraph(SUBTITLE,small=True);doc.add_paragraph(TITLE,'Title')
    lead=paragraph(LEAD,bold=True)
    for r in lead.runs:r.font.size=Pt(11)
    paragraph(SUMMARY);doc.add_paragraph('핵심 비교 결과','Heading 1')
    table=doc.add_table(rows=1,cols=4);table.autofit=False
    widths=[2.4,1.3,1.5,1.5]
    for i,h in enumerate(HEADERS):
        cell=table.rows[0].cells[i];cell.text=h;shade(cell,'173750')
        for r in cell.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    for index,row in enumerate(rows):
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=value
            if index==TABLE_HIGHLIGHT_ROW-1:shade(cells[i],'EDF6F4')
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i])
            for p in cell.paragraphs:
                p.alignment=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
    paragraph(TABLE_NOTE,small=True);paragraph(TERMS,small=True);paragraph(INSIGHT,bold=True)
    doc.add_paragraph('설계 의사결정 제안','Heading 1')
    for label,body in DECISIONS:
        p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(body)
    paragraph(GATING,small=True);doc.add_paragraph('판단 범위와 근거','Heading 1')
    for text in [SCOPE,LIMIT,VALIDATION,SOURCE]:paragraph(text,small=True)
    doc.core_properties.title=TITLE
    doc.core_properties.subject='동일 조회 병렬성의 클럭·지연·파워·에너지 비교'
    path=OUT/f'{STEM}.docx';doc.save(path)
    check=Document(path)
    assert len(check.tables)==1 and len(check.tables[0].rows)==len(rows)+1
    assert [[c.text for c in row.cells] for row in check.tables[0].rows[1:]]==rows


def main():
    rows=get_rows();build_pdf(rows);build_docx(rows)
    manifest=json.loads(MANIFEST.read_text())
    evidence=dict(source_csv=str(CSV.relative_to(ROOT)),
                  source_csv_sha256=hashlib.sha256(CSV.read_bytes()).hexdigest(),
                  source_manifest_sha256=hashlib.sha256(MANIFEST.read_bytes()).hexdigest(),
                  source_simulator_cases=manifest['simulator_cases'],source_energy_rows=manifest['energy_rows'],
                  build_script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                  verified_pdf_pages=1,table_rows=rows)
    (OUT/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
    print(OUT)


if __name__=='__main__':main()
