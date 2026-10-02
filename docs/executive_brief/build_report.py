"""Generate the executive brief from the checked-in finite-lookup results."""
import csv
from pathlib import Path
import tempfile
import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
SOURCE_COMMIT='55e06c353480aa11da70156dad4324700f773bb0'
SOURCE=f'https://github.com/JBRARAPARK/NAND_Prefetch_Model/blob/{SOURCE_COMMIT}/results/lookup_slots_64b/RESULTS_KO.md'
CSV=ROOT/'results/lookup_slots_64b/sweep.csv'
TITLE='NAND 프리페치 조회 병렬성 검토'
SUBTITLE='임원 보고  |  2026년 10월 2일  |  모델 기반 설계 검토'
LEAD='조회 처리능력이 병목이면 요청 수 확대만으로 성능은 개선되지 않습니다.'
SUMMARY='프리페치 데이터가 준비돼 있어도, 사용 가능 여부를 확인하는 조회가 지연되면 전송 대역폭을 활용하지 못합니다. 설계 우선순위는 동시 요청 수 확대보다 조회 처리능력과 요청 수의 균형 확보입니다.'
HEADERS=['동시 조회 슬롯','처리량 GB/s','평균 조회 대기 µs']
TERMS='조회 슬롯은 동시에 진행 중인 판정 수이며, outstanding은 수락 후 반환이 끝나지 않은 요청 수입니다. 조회는 500ns 동안 슬롯을 점유하고, 초과 요청은 순서대로 대기합니다.'
INSIGHT='조회 슬롯 64개에서 outstanding을 100개에서 500개로 늘려도 처리량은 8.19GB/s로 동일합니다. 평균 조회 대기는 0.25µs에서 3.34µs로 늘어납니다. 과도한 요청 확대는 처리량 개선 없이 응답 대기를 키울 수 있습니다.'
DECISIONS=[
 ('우선 검증 대상', '실제 조회 구조의 새 요청 수락 간격과 동시 처리 용량을 측정하고, 이를 기준으로 outstanding을 정합니다.'),
 ('32GB/s 목표의 조건', '조회 500ns가 유지될 때, 준비된 데이터의 연속 전송에는 조회 슬롯 250개 이상과 outstanding 251개 이상이 함께 필요합니다. 슬롯 수는 물리 메모리 포트 수가 아닙니다.'),
 ('설계 확정 전 확인', '파이프라인·메모리 충돌·실제 hit 비율과 랜덤 접근을 반영해 재검증합니다. 현재 결과만으로 제품 성능이나 회로 면적·전력 비용을 확정할 수 없습니다.')]
SCOPE='조건: tR 2.08µs, 64B AXI 읽기, 256bit·1GHz, 4칩·4채널. 64KiB 요청을 64B로 나누며 NAND page는 16KiB입니다. 표는 데이터가 모두 사전 준비된 경우로, 준비에 사용한 NAND 시간은 처리량 계산에서 제외합니다. 31.76GB/s는 초기 조회 대기를 포함한 값이며 이후 연속 전송은 32GB/s입니다.'
VALIDATION='검증: 260개 시뮬레이션 조건, 테스트 26개 통과, 기존 32개 조건과 회귀 일치. 실제 제품 실측은 수행하지 않았습니다.'


def get_rows():
    with CSV.open() as f: rows=list(csv.DictReader(f))
    out=[]
    for c in ['4','64','128','250','unlimited']:
        r=next(r for r in rows if r['scenario']=='ready' and r['outstanding']=='500' and r['lookup_slots']==c and r['lookup_us']=='0.5')
        out.append(['무제한' if c=='unlimited' else c,
                    f'{float(r["throughput_gbps"]):.2f}', f'{float(r["mean_lookup_wait_us"]):.2f}'])
    return out


def build_pdf(rows):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as ReportFont
    import urllib.request
    fontdir=Path(tempfile.gettempdir())/'nand-executive-fonts';fontdir.mkdir(exist_ok=True)
    for weight,name in [('Regular','Korean'),('Bold','KoreanBold')]:
        path=fontdir/f'NanumGothic-{weight}.ttf'
        if not path.exists():
            urllib.request.urlretrieve(f'https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/NanumGothic-{weight}.ttf',path)
        pdfmetrics.registerFont(ReportFont(name,str(path)))
    pdfmetrics.registerFontFamily('Korean',normal='Korean',bold='KoreanBold')
    navy=colors.HexColor('#173750');teal=colors.HexColor('#147B80')
    styles={
        'body':ParagraphStyle('body',fontName='Korean',fontSize=9.4,leading=14.2,textColor=navy,spaceAfter=6,wordWrap='CJK'),
        'small':ParagraphStyle('small',fontName='Korean',fontSize=8,leading=11.4,textColor=colors.HexColor('#536779'),spaceAfter=5,wordWrap='CJK'),
        'title':ParagraphStyle('title',fontName='KoreanBold',fontSize=20,leading=26,textColor=navy,spaceAfter=13),
        'head':ParagraphStyle('head',fontName='KoreanBold',fontSize=10.6,leading=15,textColor=teal,spaceBefore=8,spaceAfter=5),
        'lead':ParagraphStyle('lead',fontName='KoreanBold',fontSize=11.3,leading=17,textColor=navy,wordWrap='CJK'),
        'cell':ParagraphStyle('cell',fontName='Korean',fontSize=9.5,leading=14,alignment=1,textColor=navy),
        'th':ParagraphStyle('th',fontName='KoreanBold',fontSize=9,leading=14,alignment=1,textColor=colors.white)}
    def p(text,kind='body'):return Paragraph(text,styles[kind])
    story=[p(SUBTITLE,'small'),Spacer(1,4),p(TITLE,'title')]
    callout=Table([[p(LEAD,'lead')]],colWidths=[523])
    callout.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E8F3F4')),('BOX',(0,0),(-1,-1),0,colors.white),('LEFTPADDING',(0,0),(-1,-1),11),('RIGHTPADDING',(0,0),(-1,-1),11),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]))
    story += [callout,Spacer(1,9),p(SUMMARY),p('주요 결과','head'),p('Outstanding 500개 · 조회 지연 500ns · 사전 준비 데이터 조건','small')]
    table=Table([[p(x,'th') for x in HEADERS]]+[[p(x,'cell') for x in row] for row in rows],colWidths=[174,174,175])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),('BACKGROUND',(0,4),(-1,4),colors.HexColor('#EDF6F4')),('LINEBELOW',(0,1),(-1,-1),.4,colors.HexColor('#D7E1E9')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    story += [table,Spacer(1,7),p(TERMS,'small'),p('<b>'+INSIGHT+'</b>'),p('설계 의사결정 제안','head')]
    for label,body in DECISIONS:story.append(p('<b>'+label+'</b>　'+body))
    story += [p('판단 범위와 근거','head'),p(SCOPE,'small'),p(VALIDATION,'small'),
              p(f'<link href="{SOURCE}" color="#147B80">출처　NAND Prefetch Model 시뮬레이션 결과와 상세 가정</link>','small')]
    path=OUT/'NAND_Prefetch_Executive_Brief_KO.pdf'
    doc=SimpleDocTemplate(str(path),pagesize=(595.276,841.89),leftMargin=36,rightMargin=36,topMargin=30,bottomMargin=30,title=TITLE)
    doc.build(story)
    pdf=fitz.open(path)
    assert len(pdf)==1, f'Expected 1 page, got {len(pdf)}'
    text=pdf[0].get_text().replace('\u00a0',' ')
    for expected in ['주요 결과','설계 의사결정 제안','판단 범위와 근거','31.76','61.52','2.08','250']:
        assert expected in text,expected
    pdf[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(OUT/'preview.png')
    print('PDF: 1 page; Korean text and key figures verified')
    pdf.close()


def shade(cell,fill):
    props=cell._tc.get_or_add_tcPr()
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);props.append(x)


def build_docx(rows):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.48)
    sec.left_margin=sec.right_margin=Inches(.55)
    normal=doc.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(9)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',20),('Heading 1',10.5)]:
        st=doc.styles[name];st.font.name='맑은 고딕';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('173750')
        st._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(5)
    p=doc.add_paragraph(SUBTITLE)
    for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('687A8B')
    doc.add_paragraph(TITLE,'Title')
    p=doc.add_paragraph();r=p.add_run(LEAD);r.bold=True;r.font.size=Pt(11)
    doc.add_paragraph(SUMMARY)
    doc.add_paragraph('주요 결과','Heading 1')
    doc.add_paragraph('Outstanding 500개 · 조회 지연 500ns · 사전 준비 데이터 조건')
    table=doc.add_table(rows=1,cols=3);table.autofit=False
    widths=[2.1,2.1,2.8]
    for i,h in enumerate(HEADERS):
        cell=table.rows[0].cells[i];cell.text=h;shade(cell,'173750')
        for r in cell.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=value
            if row[0]=='250':shade(cells[i],'EDF6F4')
    for row in table.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i])
            for p in cell.paragraphs:
                p.alignment=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
    p=doc.add_paragraph(TERMS)
    for r in p.runs:r.font.size=Pt(8)
    p=doc.add_paragraph(INSIGHT)
    for r in p.runs:r.bold=True
    doc.add_paragraph('설계 의사결정 제안','Heading 1')
    for label,body in DECISIONS:
        p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(body)
    doc.add_paragraph('판단 범위와 근거','Heading 1')
    for text in [SCOPE,VALIDATION]:
        p=doc.add_paragraph(text)
        for r in p.runs:r.font.size=Pt(8)
    p=doc.add_paragraph('출처  ')
    rel=p.part.relate_to(SOURCE,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');txt=OxmlElement('w:t');txt.text='NAND Prefetch Model 시뮬레이션 결과와 상세 가정';run.append(txt);link.append(run);p._p.append(link)
    doc.core_properties.title=TITLE
    doc.core_properties.subject='조회 병렬성 제한에 따른 성능과 설계 의사결정'
    path=OUT/'NAND_Prefetch_Executive_Brief_KO.docx';doc.save(path)
    check=Document(path);assert len(check.tables)==1 and len(check.tables[0].rows)==6


if __name__=='__main__':
    rows=get_rows();build_pdf(rows);build_docx(rows)
    print(OUT)
